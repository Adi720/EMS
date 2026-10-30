import io
from datetime import date, timedelta

import pytest
from django.db.models import ProtectedError
from django.urls import reverse
from openpyxl import load_workbook

from apps.allocations.services import current_allocation
from apps.audit.models import AuditEvent
from apps.calling.models import TaskStatus
from apps.contacts.models import Contact, ContactStatus, ContactStatusChange, StatusReason
from apps.contacts.services import LifecycleError, change_status

VALID = {
    "full_name": "  Asha   Synthetic ",
    "phone_primary": "98765 43210",
    "phone_secondary": "",
    "area": "Riverside",
    "registration_date": "2026-03-01",
    "qualification": "",
    "collected_by_name": "",
    "notes": "",
}


@pytest.fixture
def coord_client(client, coordinator):
    client.force_login(coordinator)
    return client


def test_create_contact_normalises_and_records_history(coord_client):
    response = coord_client.post(reverse("contacts:create"), VALID)
    assert response.status_code == 302
    contact = Contact.objects.get()
    assert contact.full_name == "Asha Synthetic"
    assert contact.phone_primary == "+919876543210"
    assert contact.status == ContactStatus.ACTIVE
    history = contact.status_history.get()
    assert (history.from_status, history.to_status, history.reason_category) == (
        "",
        "active",
        StatusReason.CREATED,
    )
    assert AuditEvent.objects.filter(action="contact.create", object_id=str(contact.pk)).exists()


def test_create_with_caller_uses_allocation_date_not_registration_date(coord_client, caller):
    data = {**VALID, "allocate_to": caller.pk, "allocation_start": "2026-10-01"}
    coord_client.post(reverse("contacts:create"), data)
    allocation = current_allocation(Contact.objects.get())
    assert allocation.caller == caller
    assert allocation.effective_from == date(2026, 10, 1)


@pytest.mark.parametrize(
    ("field", "value"),
    [("phone_primary", "12345"), ("full_name", ""), ("area", " "), ("registration_date", "")],
)
def test_invalid_input_is_rejected(coord_client, field, value):
    response = coord_client.post(reverse("contacts:create"), {**VALID, field: value})
    assert response.status_code == 200
    assert response.context["form"].errors[field]
    assert not Contact.objects.exists()


def test_future_registration_date_is_rejected(coord_client):
    future = (date.today() + timedelta(days=5)).isoformat()
    response = coord_client.post(reverse("contacts:create"), {**VALID, "registration_date": future})
    assert "registration_date" in response.context["form"].errors


def test_duplicate_phone_warns_and_requires_confirmation(coord_client, make_contact):
    existing = make_contact(full_name="Existing Person", phone="+919876543210")
    response = coord_client.post(reverse("contacts:create"), VALID)
    assert response.status_code == 200
    assert existing in response.context["form"].duplicates
    assert Contact.objects.count() == 1
    response = coord_client.post(reverse("contacts:create"), {**VALID, "confirm_duplicate": "on"})
    assert response.status_code == 302
    assert Contact.objects.count() == 2  # never merged automatically


def test_edit_contact(coord_client, make_contact):
    contact = make_contact(full_name="Old Name")
    data = {**VALID, "full_name": "New Name", "phone_primary": contact.phone_display}
    response = coord_client.post(reverse("contacts:edit", args=[contact.pk]), data)
    assert response.status_code == 302
    contact.refresh_from_db()
    assert contact.full_name == "New Name"
    event = AuditEvent.objects.get(action="contact.update")
    assert "full_name" in event.details["fields"]
    assert contact.phone_primary not in str(event.details)  # no phone numbers in the audit log


def test_list_search_filters_and_paginates(coord_client, make_contact, caller, allocate):
    target = make_contact(full_name="Findable Person", phone="+919123456789", area="Hill View")
    allocate(target, caller)
    for _ in range(30):
        make_contact()
    assert coord_client.get(reverse("contacts:list"), {"q": "23456"}).context["object_list"][0] == target
    assert list(coord_client.get(reverse("contacts:list"), {"area": "hill"}).context["object_list"]) == [
        target
    ]
    assert list(coord_client.get(reverse("contacts:list"), {"caller": caller.pk}).context["object_list"]) == [
        target
    ]
    unallocated = coord_client.get(reverse("contacts:list"), {"caller": "none"})
    assert unallocated.context["paginator"].count == 30
    assert len(unallocated.context["object_list"]) == 25


# --- Lifecycle -------------------------------------------------------------------


def test_temporarily_unavailable_keeps_allocation_and_withdraws_unattempted(
    make_contact, caller, allocate, make_session, coordinator
):
    contact = make_contact()
    allocate(contact, caller)
    session = make_session()
    change_status(
        contact,
        ContactStatus.TEMP_UNAVAILABLE,
        coordinator,
        reason_category=StatusReason.TEMPORARILY_AWAY,
        review_date=date(2026, 12, 1),
    )
    contact.refresh_from_db()
    assert contact.review_date == date(2026, 12, 1)
    assert current_allocation(contact).caller == caller
    task = session.tasks.get(contact=contact)
    assert task.status == TaskStatus.WITHDRAWN
    change_status(contact, ContactStatus.ACTIVE, coordinator, reason_category=StatusReason.AVAILABLE_AGAIN)
    task.refresh_from_db()
    assert task.status == TaskStatus.OPEN and task.caller == caller


@pytest.mark.parametrize("status", [ContactStatus.INACTIVE, ContactStatus.DO_NOT_CALL])
def test_inactive_and_do_not_call_end_allocation_keep_history(
    make_contact, caller, allocate, coordinator, status
):
    contact = make_contact()
    allocation = allocate(contact, caller).opened.get()
    change_status(
        contact,
        status,
        coordinator,
        reason_category=StatusReason.MOVED_AWAY,
        reason_text="Moved to another city",
        effective_date=date(2026, 6, 1),
    )
    allocation.refresh_from_db()
    assert allocation.effective_to == date(2026, 6, 1)
    assert current_allocation(contact) is None
    assert contact.allocations.count() == 1
    history = contact.status_history.first()
    assert (history.from_status, history.to_status, history.reason_text) == (
        "active",
        status,
        "Moved to another city",
    )
    assert history.changed_by == coordinator


def test_leaving_do_not_call_requires_consent_note(make_contact, coordinator, caller):
    contact = make_contact()
    change_status(
        contact, ContactStatus.DO_NOT_CALL, coordinator, reason_category=StatusReason.REQUESTED_NO_CALLS
    )
    with pytest.raises(LifecycleError):
        change_status(
            contact, ContactStatus.ACTIVE, coordinator, reason_category=StatusReason.AVAILABLE_AGAIN
        )
    change_status(
        contact,
        ContactStatus.ACTIVE,
        coordinator,
        reason_category=StatusReason.AVAILABLE_AGAIN,
        reason_text="Called us and asked to be invited again",
        reactivation_caller=caller,
    )
    assert current_allocation(contact).caller == caller


def test_same_status_is_rejected(make_contact, coordinator):
    with pytest.raises(LifecycleError):
        change_status(make_contact(), ContactStatus.ACTIVE, coordinator, reason_category=StatusReason.OTHER)


def test_status_view_and_validation(coord_client, make_contact):
    contact = make_contact()
    url = reverse("contacts:status", args=[contact.pk])
    response = coord_client.post(
        url,
        {
            "new_status": "inactive",
            "reason_category": "moved_away",
            "effective_date": "2026-06-01",
            "review_date": "2026-07-01",
        },
    )
    assert response.status_code == 200  # review date only for temporarily unavailable
    contact.refresh_from_db()
    assert contact.status == ContactStatus.ACTIVE
    response = coord_client.post(
        url, {"new_status": "inactive", "reason_category": "moved_away", "effective_date": "2026-06-01"}
    )
    assert response.status_code == 302
    contact.refresh_from_db()
    assert contact.status == ContactStatus.INACTIVE


def test_history_is_append_only_and_contacts_cannot_be_deleted(make_contact):
    contact = make_contact()
    history = contact.status_history.get()
    history.reason_text = "tampered"
    with pytest.raises(ValueError):
        history.save()
    with pytest.raises(ValueError):
        history.delete()
    with pytest.raises(ProtectedError):
        contact.delete()
    assert ContactStatusChange.objects.filter(pk=history.pk, reason_text="").exists()


# --- Export ------------------------------------------------------------------------


def test_export_excludes_do_not_call_and_neutralises_formulas(coord_client, make_contact, coordinator):
    make_contact(full_name='=HYPERLINK("http://evil.example","x")')
    dnc = make_contact(full_name="Private Person")
    change_status(
        dnc, ContactStatus.DO_NOT_CALL, coordinator, reason_category=StatusReason.REQUESTED_NO_CALLS
    )

    response = coord_client.get(reverse("contacts:export"), {"format": "csv"})
    body = response.content.decode("utf-8-sig")
    assert "Private Person" not in body
    assert "'=HYPERLINK" in body

    response = coord_client.get(reverse("contacts:export"), {"format": "xlsx", "include_dnc": "1"})
    sheet = load_workbook(io.BytesIO(response.content)).active
    names = [row[0].value for row in sheet.iter_rows(min_row=2)]
    assert "Private Person" in names
    assert all(not str(n).startswith("=") for n in names)
    assert sheet.cell(row=2, column=1).data_type == "s"
    assert AuditEvent.objects.filter(action="export.contacts").count() == 2
