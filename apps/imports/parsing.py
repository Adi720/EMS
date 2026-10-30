"""Reading XLSX/CSV uploads into plain rows. No database access here."""

import csv
import hashlib
import io
import re
from datetime import date, datetime, time, timedelta

from django.conf import settings
from openpyxl import load_workbook

from apps.core.text import clean_text

CONTACT_SHEET = "ems stall sewa"
STATUS_SHEET = "2026 calling sewa status"
MAX_COLUMNS = 80
EXCEL_EPOCH = date(1899, 12, 30)


class ParseError(Exception):
    pass


# Import targets for the contact sheet: (key, label, required, header aliases).
TARGET_FIELDS = [
    ("registration_date", "Registration date (stall date)", True, ["date"]),
    ("full_name", "Name", True, ["name"]),
    (
        "phone",
        "Mobile number",
        True,
        ["mobile no", "mobile no.", "mobile", "mobile number", "phone", "contact number"],
    ),
    ("area", "Area", True, ["area"]),
    ("qualification", "Qualification", False, ["qualification"]),
    ("collected_by", "Collected by (stall seva member)", False, ["sewa conducted by", "seva conducted by"]),
    (
        "caller",
        "Allocated caller",
        False,
        [
            "calling co-ordinater",
            "calling co-ordinator",
            "calling coordinator",
            "calling coordinater",
            "caller",
        ],
    ),
]
TARGET_LABELS = {key: label for key, label, _, _ in TARGET_FIELDS}
REQUIRED_TARGETS = [key for key, _, required, _ in TARGET_FIELDS if required]

STATUS_FIELDS = {
    "name": ["name"],
    "contact_number": ["contact number", "contact no", "mobile no"],
    "total_assigned": ["total assigned"],
    "coordinator": ["co-ordinator", "coordinator", "co-ordinater"],
}


def norm_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def sha256_of(uploaded_file):
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    return digest.hexdigest()


def cell_to_json(value):
    """JSON-safe representation that keeps numbers and dates distinguishable from text."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return value
    text = str(value).strip()
    return text or None


def read_sheets(uploaded_file):
    """Return (file_type, {sheet name: [row tuples]})."""
    name = uploaded_file.name.lower()
    if uploaded_file.size > settings.IMPORT_MAX_UPLOAD_BYTES:
        raise ParseError("The file is too large.")
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise ParseError("The file could not be read as an Excel (.xlsx) workbook.") from exc
        sheets = {}
        try:
            for sheet in workbook.worksheets:
                rows = []
                for index, row in enumerate(sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)):
                    if index > settings.IMPORT_MAX_ROWS + 20:
                        raise ParseError(
                            f"Sheet '{sheet.title}' has more than {settings.IMPORT_MAX_ROWS} rows."
                        )
                    rows.append(tuple(cell_to_json(v) for v in row))
                sheets[sheet.title] = rows
        finally:
            workbook.close()
        return "xlsx", sheets
    if name.endswith(".csv"):
        raw = uploaded_file.read()
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ParseError("The CSV file encoding is not supported.")
        rows = [tuple(cell_to_json(v) for v in row[:MAX_COLUMNS]) for row in csv.reader(io.StringIO(text))]
        if len(rows) > settings.IMPORT_MAX_ROWS + 20:
            raise ParseError(f"The file has more than {settings.IMPORT_MAX_ROWS} rows.")
        return "csv", {"CSV": rows}
    raise ParseError("Upload an .xlsx or .csv file.")


def find_sheet(sheets, wanted):
    for name in sheets:
        if norm_header(name) == wanted:
            return name
    return None


def detect_header_row(rows, known_headers, scan=10):
    """Index of the first row containing at least two known header names (else 0)."""
    known = {norm_header(h) for h in known_headers}
    for index, row in enumerate(rows[:scan]):
        if sum(1 for cell in row if norm_header(cell) in known) >= 2:
            return index
    return 0


def table_from_rows(rows, header_index):
    """Headers (unique, cleaned) and [(sheet row number, {header: value})] for non-empty rows."""
    header_cells = rows[header_index] if rows else ()
    headers, seen = [], {}
    for position, cell in enumerate(header_cells):
        label = clean_text(cell) or f"Column {position + 1}"
        count = seen.get(label, 0) + 1
        seen[label] = count
        headers.append(label if count == 1 else f"{label} ({count})")
    records = []
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if row[i] is not None}
        if values:
            records.append((offset, values))
    return headers, records


def guess_mapping(headers):
    by_norm = {norm_header(h): h for h in headers}
    mapping = {}
    for key, _, _, aliases in TARGET_FIELDS:
        for alias in aliases:
            if alias in by_norm:
                mapping[key] = by_norm[alias]
                break
    return mapping


def parse_status_sheet(rows):
    """Caller records from '2026 Calling Sewa status' (requirements B-2 interpretation)."""
    known = [alias for aliases in STATUS_FIELDS.values() for alias in aliases]
    header_index = detect_header_row(rows, known)
    headers, records = table_from_rows(rows, header_index)
    columns = {}
    by_norm = {norm_header(h): h for h in headers}
    for key, aliases in STATUS_FIELDS.items():
        columns[key] = next((by_norm[a] for a in aliases if a in by_norm), None)
    if not columns["name"]:
        return [], headers
    result = []
    for row_number, values in records:
        name = clean_text(values.get(columns["name"]))
        if not name or norm_header(name) in {"total", "grand total"}:
            continue
        total_raw = values.get(columns["total_assigned"]) if columns["total_assigned"] else None
        total = parse_int(total_raw)
        result.append(
            {
                "row": row_number,
                "name": name,
                "contact_number": str(values.get(columns["contact_number"]) or "")
                if columns["contact_number"]
                else "",
                "total_assigned": total,
                "total_assigned_raw": "" if total_raw is None else str(total_raw),
                "coordinator": clean_text(values.get(columns["coordinator"]))
                if columns["coordinator"]
                else "",
            }
        )
    return result, headers


def parse_int(value):
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


DATE_FORMATS = [
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%d %b %Y",
    "%d %B %Y",
    "%d-%b-%Y",
]


def parse_date_cell(value):
    """Return (date | None, ambiguous: bool). Text dates are read day-first (A-1)."""
    if value is None or value == "":
        return None, False
    if isinstance(value, int | float):
        if 20000 < value < 80000:  # plausible Excel serial date (1954–2119)
            return EXCEL_EPOCH + timedelta(days=int(value)), False
        return None, False
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10]), False
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        parts = re.split(r"[/\-. ]", text)
        ambiguous = (
            len(parts) >= 2
            and parts[0].isdigit()
            and parts[1].isdigit()
            and int(parts[0]) <= 12
            and int(parts[1]) <= 12
            and parts[0] != parts[1]
        )
        return parsed, ambiguous
    return None, False
