"""Synthetic workbook in the layout of '2026 EMS DATA.xlsx' (docs/import-mapping.md §7).

All names and numbers are generated. The data deliberately contains the problems
the importer must handle; ``SYNTHETIC_FACTS`` records which rows have which problem.
"""

import io
import random
from datetime import date, timedelta

import phonenumbers
from openpyxl import Workbook

FIRST = [
    "Aarav",
    "Vihaan",
    "Ishaan",
    "Kabir",
    "Reyansh",
    "Anaya",
    "Diya",
    "Myra",
    "Sara",
    "Kiara",
    "Arjun",
    "Rohan",
    "Karan",
    "Nikhil",
    "Tanvi",
    "Riya",
    "Meera",
    "Simran",
    "Aditya",
    "Varun",
]
LAST = ["Synthetic", "Example", "Testwala", "Demo", "Sample", "Placeholder", "Mockkar", "Fakeja"]
AREAS = ["North Block", "Riverside", "Old Town", "Hill View", "Market Road", "Station Area"]
CALLERS = ["Synth Caller One", "Synth Caller Two", "Synth Caller Three", "Synth Caller Four"]
COORDINATORS = ["Synth Coordinator A", "Synth Coordinator B"]
COLLECTORS = ["Stall Member X", "Stall Member Y", "Stall Member Z"]

CONTACT_HEADERS = [
    "DATE",
    "NAME",
    "MOBILE NO",
    "AREA",
    "QUALIFICATION",
    "SEWA CONDUCTED BY",
    "CALLING CO-ORDINATER",
    "WEEK 1",
    "WEEK 2",
    "WEEK 3",
]

# Excel rows (header is row 1) carrying deliberate problems.
SYNTHETIC_FACTS = {
    "duplicate_of_previous_row": 6,  # same phone + name as row 5
    "shared_phone_different_name": 9,  # same phone as row 8, different person
    "invalid_phone": 11,
    "two_phones": 13,
    "blank_area": 17,
    "text_date": 19,
    "serial_date": 20,
    "caller_spelling_variant": 21,  # "  synth caller one " → Synth Caller One
    "unknown_caller": 23,
    "blank_caller": 25,
    "float_phone": 27,
}


def _phone(rng, used):
    while True:
        candidate = f"9{rng.randint(100000000, 999999999)}"
        if candidate in used:
            continue
        if phonenumbers.is_valid_number(phonenumbers.parse(candidate, "IN")):
            used.add(candidate)
            return candidate


MIN_CONTACTS = 30  # enough rows for every deliberate problem in SYNTHETIC_FACTS


def build_rows(n_contacts=120, seed=2026):
    if n_contacts < MIN_CONTACTS:
        raise ValueError(f"The synthetic workbook needs at least {MIN_CONTACTS} contacts.")
    rng = random.Random(seed)
    used = set()
    rows = []
    start = date(2026, 1, 4)
    for index in range(n_contacts):
        excel_row = index + 2
        rows.append(
            {
                "DATE": start + timedelta(days=7 * (index // 10)),
                "NAME": f"{FIRST[index % len(FIRST)]} {LAST[(index // len(FIRST)) % len(LAST)]} {index + 1}",
                "MOBILE NO": int(_phone(rng, used)),
                "AREA": AREAS[index % len(AREAS)],
                "QUALIFICATION": rng.choice(["", "B.Com", "BSc", "12th", "MBA"]),
                "SEWA CONDUCTED BY": COLLECTORS[index % len(COLLECTORS)],
                "CALLING CO-ORDINATER": CALLERS[index % len(CALLERS)],
                "WEEK 1": rng.choice(["Yes", "No", "", "NC"]),
                "WEEK 2": rng.choice(["Yes", "No", "", "NC"]),
                "WEEK 3": rng.choice(["", "Yes"]),
                "_excel_row": excel_row,
            }
        )
    by_row = {r["_excel_row"]: r for r in rows}
    facts = SYNTHETIC_FACTS
    dup = by_row[facts["duplicate_of_previous_row"]]
    prev = by_row[facts["duplicate_of_previous_row"] - 1]
    dup.update({"NAME": prev["NAME"], "MOBILE NO": prev["MOBILE NO"]})
    by_row[facts["shared_phone_different_name"]]["MOBILE NO"] = by_row[
        facts["shared_phone_different_name"] - 1
    ]["MOBILE NO"]
    by_row[facts["invalid_phone"]]["MOBILE NO"] = "12345"
    by_row[facts["two_phones"]]["MOBILE NO"] = f"{_phone(rng, used)} / {_phone(rng, used)}"
    by_row[facts["blank_area"]]["AREA"] = None
    by_row[facts["text_date"]]["DATE"] = "15/03/2026"
    by_row[facts["serial_date"]]["DATE"] = 46100  # an Excel serial date (March 2026)
    by_row[facts["caller_spelling_variant"]]["CALLING CO-ORDINATER"] = "  synth caller one "
    by_row[facts["unknown_caller"]]["CALLING CO-ORDINATER"] = "Unknown Person"
    by_row[facts["blank_caller"]]["CALLING CO-ORDINATER"] = None
    by_row[facts["float_phone"]]["MOBILE NO"] = float(_phone(rng, used))
    return rows


def expected_caller_counts(rows):
    counts = {}
    for row in rows:
        name = " ".join(str(row["CALLING CO-ORDINATER"] or "").split()).lower()
        counts[name] = counts.get(name, 0) + 1
    return counts


def build_workbook(n_contacts=120, seed=2026, *, mismatch_caller=CALLERS[1]):
    """Return (xlsx bytes, rows). One caller's Total Assigned is deliberately off by one."""
    rows = build_rows(n_contacts, seed)
    workbook = Workbook()
    contacts = workbook.active
    contacts.title = "EMS STALL SEWA"
    contacts.append(CONTACT_HEADERS)
    for row in rows:
        contacts.append([row[h] for h in CONTACT_HEADERS])

    status = workbook.create_sheet("2026 Calling Sewa status")
    status.append(
        ["Name", "Contact number", "Total Assigned", "Co-ordinator", "05 Jan Attempted", "05 Jan Confirmed"]
    )
    counts = expected_caller_counts(rows)
    for index, caller in enumerate(CALLERS):
        total = sum(n for name, n in counts.items() if name == caller.lower())
        if caller == mismatch_caller:
            total += 1
        status.append([caller, f"98000000{index:02d}", total, COORDINATORS[index % 2], 5, 2])
    status.append(["Old Synth Caller", "9800000099", 0, COORDINATORS[0], 0, 0])

    weekly = workbook.create_sheet("Weekly Data")
    weekly.append(["Date", "Callers", "Calls to be Made", "Calls Made", "Confirmed Saints", "Saints in EMS"])
    weekly.append([date(2026, 1, 4), 4, 120, 100, 30, 25])
    workbook.create_sheet("Calling Sewa status").append(["Name", "Total"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), rows
