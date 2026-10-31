from datetime import date

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.accounts.models import User
from apps.accounts.roles import Role
from apps.allocations.models import Allocation, AllocationCycle, ChangeType
from apps.allocations.services import create_cycle
from apps.audit.models import AuditEvent
from apps.contacts.models import Contact
from apps.imports import parsing, services
from apps.imports.models import BatchStatus, RowAction, RowStatus
from apps.imports.synthetic import CALLERS, SYNTHETIC_FACTS, build_workbook

CYCLE_START = date(2026, 9, 1)


@pytest.fixture
def synthetic_callers(make_user):
    users = {}
    for index, name in enumerate(CALLERS):
        first, last = name.split(" ", 1)
        users[name] = make_user(f"synth{index}", [Role.CALLER], first_name=first, last_name=last)
    return users


@pytest.fixture
def workbook():
    content, rows = build_workbook()
    return content, rows


def upload(content, name="2026 EMS DATA (synthetic).xlsx"):
    return SimpleUploadedFile(name, content, content_type="application/octet-stream")


@pytest.fixture
def staged(workbook, coordinator, synthetic_callers):
    batch = services.stage_upload(upload(workbook[0]), coordinator, use_status_sheet=True)
    batch.mapping_confirmed = True
    batch.save()
    services.validate_batch(batch)
    batch.refresh_from_db()
    return batch


def row(batch, number):
    return batch.rows.get(row_number=number)


def codes(r):
    return {i["code"] for i in r.issues}


def resolve_everything(batch, coordinator, synthetic_callers):
    services.set_caller_mapping(batch, "Unknown Person", str(synthetic_callers[CALLERS[0]].pk), coordinator)
    services.set_caller_mapping(batch, "", services.UNALLOCATED, coordinator)
    for r in batch.rows.filter(status=RowStatus.NEEDS_RESOLUTION):
        if "missing_area" in codes(r):
            services.update_row(
                r,
                overrides={"area": "Riverside"},
                action="",
                link_contact=None,
                skip_reason="",
                actor=coordinator,
            )
        else:
            services.update_row(
                r,
                overrides={},
                action=RowAction.SKIP,
                link_contact=None,
                skip_reason="Duplicate / invalid in source",
                actor=coordinator,
            )
    batch.refresh_from_db()


def prepare_commit(batch, coordinator):
    batch.cycle = create_cycle(
        name="Sep–Nov 2026", start_date=CYCLE_START, planned_end_date=date(2026, 11, 30), actor=coordinator
    )
    batch.caller_column_confirmed = True
    batch.discrepancies_acknowledged = True
    batch.save()


def test_staging_reads_sheets_mapping_and_caller_records(staged):
    assert staged.contact_sheet == "EMS STALL SEWA"
    assert staged.status_sheet == "2026 Calling Sewa status"
    assert staged.rows.count() == 120
    assert staged.column_mapping == {
        "registration_date": "DATE",
        "full_name": "NAME",
        "phone": "MOBILE NO",
        "area": "AREA",
        "qualification": "QUALIFICATION",
        "collected_by": "SEWA CONDUCTED BY",
        "caller": "CALLING CO-ORDINATER",
    }
    assert [r["name"] for r in staged.caller_records] == [*CALLERS, "Old Synth Caller"]
    assert "WEEK 1" in row(staged, 2).raw_values  # historical columns kept in staging
    assert not Contact.objects.exists()  # nothing imported yet


def test_row_validation_flags_each_synthetic_problem(staged, synthetic_callers):
    f = SYNTHETIC_FACTS
    assert "duplicate_in_file" in codes(row(staged, f["duplicate_of_previous_row"]))
    assert row(staged, f["shared_phone_different_name"]).status == RowStatus.WARNING
    assert "invalid_phone" in codes(row(staged, f["invalid_phone"]))
    two = row(staged, f["two_phones"])
    assert "two_phones" in codes(two) and two.values["phone_secondary"]
    assert "missing_area" in codes(row(staged, f["blank_area"]))
    assert row(staged, f["text_date"]).values["registration_date"] == "2026-03-15"
    assert row(staged, f["serial_date"]).values["registration_date"].startswith("2026-03")
    variant = row(staged, f["caller_spelling_variant"])
    assert variant.values["caller_id"] == synthetic_callers[CALLERS[0]].pk
    assert "caller_unmatched" in codes(row(staged, f["unknown_caller"]))
    assert "caller_blank" in codes(row(staged, f["blank_caller"]))
    floaty = row(staged, f["float_phone"])
    assert floaty.status == RowStatus.READY and floaty.values["phone"].startswith("+919")
    assert row(staged, 2).status == RowStatus.READY


def test_discrepancies_report_total_assigned_mismatch(staged):
    messages = " ".join(d["message"] for d in staged.discrepancies)
    assert "Synth Caller Two" in messages and "Total Assigned says" in messages
    assert {d["code"] for d in staged.discrepancies} >= {"count_mismatch"}
    assert "caller_column_role" not in {d["code"] for d in staged.discrepancies}


def test_commit_is_blocked_until_everything_is_resolved_and_confirmed(staged, coordinator, synthetic_callers):
    with pytest.raises(services.ImportProblem):
        services.commit_batch(staged, coordinator)
    resolve_everything(staged, coordinator, synthetic_callers)
    assert not staged.rows.filter(status=RowStatus.NEEDS_RESOLUTION).exists()
    with pytest.raises(services.ImportProblem) as exc:
        services.commit_batch(staged, coordinator)
    assert "caller column" in str(exc.value) and "cycle" in str(exc.value)
    assert not Contact.objects.exists()


def test_full_import_creates_contacts_and_preserves_allocations(
    staged, coordinator, synthetic_callers, workbook
):
    resolve_everything(staged, coordinator, synthetic_callers)
    prepare_commit(staged, coordinator)
    batch = services.commit_batch(staged, coordinator)
    result = batch.summary["result"]
    skipped = batch.rows.filter(status=RowStatus.SKIPPED).count()
    assert batch.status == BatchStatus.COMMITTED
    assert result["created"] == 120 - skipped == Contact.objects.count()
    assert result["unallocated"] == 1  # the blank-caller row
    allocations = Allocation.objects.all()
    assert allocations.count() == result["allocated"]
    assert {a.effective_from for a in allocations} == {CYCLE_START}  # never the registration date
    assert {a.change_type for a in allocations} == {ChangeType.INITIAL_IMPORT}
    assert {a.cycle for a in allocations} == {AllocationCycle.objects.get()}

    imported = Contact.objects.get(full_name=workbook[1][0]["NAME"])
    assert imported.registration_date == workbook[1][0]["DATE"]
    assert imported.legacy_data["original"]["MOBILE NO"] == workbook[1][0]["MOBILE NO"]
    assert imported.legacy_data["import"] == {"batch": batch.pk, "sheet": "EMS STALL SEWA", "row": 2}
    assert imported.allocations.get().caller == synthetic_callers[CALLERS[0]]
    assert AuditEvent.objects.filter(action="import.commit").exists()


def test_repeat_import_creates_nothing(staged, coordinator, synthetic_callers, workbook):
    resolve_everything(staged, coordinator, synthetic_callers)
    prepare_commit(staged, coordinator)
    services.commit_batch(staged, coordinator)
    contacts, allocations = Contact.objects.count(), Allocation.objects.count()

    again = services.stage_upload(upload(workbook[0]), coordinator)
    again.mapping_confirmed = True
    again.save()
    services.validate_batch(again)
    again.refresh_from_db()
    assert "same_file" in {d["code"] for d in again.discrepancies}
    # Every created contact matches again (plus the in-file duplicate of one of them);
    # only the invalid-phone row, which was skipped last time, still needs a decision.
    assert again.rows.filter(status=RowStatus.ALREADY_IMPORTED).count() == contacts + 1
    assert list(
        again.rows.filter(status=RowStatus.NEEDS_RESOLUTION).values_list("row_number", flat=True)
    ) == [SYNTHETIC_FACTS["invalid_phone"]]
    resolve_everything(again, coordinator, synthetic_callers)
    again.cycle = AllocationCycle.objects.get()
    again.caller_column_confirmed = True
    again.discrepancies_acknowledged = True
    again.save()
    services.commit_batch(again, coordinator)
    assert Contact.objects.count() == contacts
    assert Allocation.objects.count() == allocations


def test_commit_rolls_back_on_failure(staged, coordinator, synthetic_callers, monkeypatch):
    resolve_everything(staged, coordinator, synthetic_callers)
    prepare_commit(staged, coordinator)

    def boom(*args, **kwargs):
        raise RuntimeError("simulated failure while allocating")

    monkeypatch.setattr(services, "apply_plan", boom)
    with pytest.raises(RuntimeError):
        services.commit_batch(staged, coordinator)
    staged.refresh_from_db()
    assert staged.status == BatchStatus.STAGED
    assert not Contact.objects.exists() and not Allocation.objects.exists()


def test_caller_column_mapped_to_stall_column_is_flagged(workbook, coordinator, synthetic_callers):
    batch = services.stage_upload(upload(workbook[0]), coordinator)
    batch.column_mapping = {
        **batch.column_mapping,
        "caller": "SEWA CONDUCTED BY",
        "collected_by": "SEWA CONDUCTED BY",
    }
    batch.mapping_confirmed = True
    batch.save()
    services.validate_batch(batch)
    batch.refresh_from_db()
    assert "caller_vs_collector" in {d["code"] for d in batch.discrepancies}


def test_link_to_existing_contact_does_not_change_its_allocation(
    staged, coordinator, synthetic_callers, make_contact, allocate, caller
):
    first = row(staged, 2)
    existing = make_contact(full_name=first.values["full_name"] + " Jr", phone=first.values["phone"])
    allocate(existing, caller)
    services.validate_batch(staged)
    first.refresh_from_db()
    assert "duplicate_in_database" in codes(first)
    services.update_row(
        first, overrides={}, action=RowAction.LINK, link_contact=existing, skip_reason="", actor=coordinator
    )
    resolve_everything(staged, coordinator, synthetic_callers)
    prepare_commit(staged, coordinator)
    batch = services.commit_batch(staged, coordinator)
    assert existing.allocations.get().caller == caller
    assert any("Not changed" in n for n in batch.summary["result"]["allocation_notes"])


def test_csv_upload_and_parsing_helpers(coordinator):
    content = (
        "NAME,MOBILE NO,AREA,DATE,CALLING CO-ORDINATER\nCsv Person,+91 98765 43210,Old Town,05/01/2026,\n"
    )
    batch = services.stage_upload(upload(content.encode(), "contacts.csv"), coordinator)
    assert batch.file_type == "csv" and batch.rows.count() == 1
    assert parsing.parse_date_cell("05/01/2026") == (date(2026, 1, 5), True)
    assert parsing.parse_date_cell("25/12/2025") == (date(2025, 12, 25), False)
    assert parsing.parse_date_cell(46100)[0].year == 2026
    assert parsing.parse_date_cell("someday") == (None, False)


def test_import_views_are_coordinator_only_and_flow_works(
    client, caller, coordinator, workbook, synthetic_callers
):
    client.force_login(caller)
    assert client.get(reverse("imports:upload")).status_code == 403
    assert client.post(reverse("imports:upload"), {"file": upload(workbook[0])}).status_code == 403

    client.force_login(coordinator)
    response = client.post(reverse("imports:upload"), {"file": upload(workbook[0])})
    assert response.status_code == 302
    batch = services.ImportBatch.objects.get()
    response = client.post(reverse("imports:mapping", args=[batch.pk]), batch.column_mapping)
    assert response.status_code == 302
    response = client.get(reverse("imports:detail", args=[batch.pk]))
    assert response.status_code == 200 and response.context["blockers"]
    response = client.post(
        reverse("imports:commit", args=[batch.pk]),
        {
            "new_cycle_name": "Now",
            "new_cycle_start": "2026-09-01",
            "confirm_cycle_start": "on",
            "confirm_caller_column": "on",
            "acknowledge_discrepancies": "on",
        },
    )
    assert response.status_code == 302
    assert not Contact.objects.exists()  # unresolved rows block the commit; nothing imported
    assert not AllocationCycle.objects.exists()  # the cycle creation was rolled back too


def test_rejects_non_spreadsheet_upload(client, coordinator):
    client.force_login(coordinator)
    response = client.post(reverse("imports:upload"), {"file": SimpleUploadedFile("x.exe", b"MZ")})
    assert response.status_code == 200 and response.context["form"].errors
    response = client.post(reverse("imports:upload"), {"file": SimpleUploadedFile("x.xlsx", b"not a zip")})
    assert response.status_code == 200 and "could not be read" in str(response.context["form"].errors)


# --- Optional status sheet ---------------------------------------------------------


def test_status_sheet_is_ignored_unless_asked_for(workbook, coordinator, synthetic_callers):
    batch = services.stage_upload(upload(workbook[0]), coordinator)
    assert batch.status_sheet == "" and batch.caller_records == []
    batch.mapping_confirmed = True
    batch.save()
    services.validate_batch(batch)
    batch.refresh_from_db()
    codes_found = {d["code"] for d in batch.discrepancies}
    assert not codes_found & {
        "count_mismatch",
        "no_status_sheet",
        "status_caller_unmatched",
        "caller_not_in_status",
    }


def test_asking_for_a_missing_status_sheet_is_a_clear_error(coordinator):
    content = "NAME,MOBILE NO,AREA,DATE\nCsv Person,9876543210,Old Town,05/01/2026\n"
    with pytest.raises(services.ImportProblem, match="status"):
        services.stage_upload(upload(content.encode(), "contacts.csv"), coordinator, use_status_sheet=True)


# --- Caller accounts from the workbook ---------------------------------------------


@pytest.fixture
def staged_without_callers(workbook, coordinator):
    batch = services.stage_upload(upload(workbook[0]), coordinator)
    batch.mapping_confirmed = True
    batch.save()
    services.validate_batch(batch)
    return batch


def test_unmatched_callers_are_listed_with_usernames(staged_without_callers, make_user):
    make_user("synth.existing", ["Caller"], first_name="Synth Caller", last_name="Four")  # matches one name
    items = {u["name"]: u for u in services.unmatched_callers(staged_without_callers)}
    assert set(items) == {"Synth Caller One", "Synth Caller Two", "Synth Caller Three", "Unknown Person"}
    assert items["Synth Caller One"]["username"] == "synth.caller.one"
    assert len({u["username"] for u in items.values()}) == 4


def test_create_callers_from_workbook(client, coordinator, staged_without_callers):
    batch = staged_without_callers
    items = services.unmatched_callers(batch)
    data = {"create": [], **{f"name_{i}": u["name"] for i, u in enumerate(items)}}
    for i, u in enumerate(items):
        data[f"username_{i}"] = u["username"]
        if u["name"].startswith("Synth"):
            data["create"].append(str(i))
    client.force_login(coordinator)
    response = client.post(reverse("imports:create_callers", args=[batch.pk]), data)
    assert response.status_code == 200
    assert "no-store" in response["Cache-Control"]
    created = response.context["created"]
    assert len(created) == 4  # "Synth Caller One" also covers its spelling variant row
    for _name, username, password in created:
        user = User.objects.get(username=username)
        assert set(user.groups.values_list("name", flat=True)) == {"Caller"}
        assert not user.is_staff and not user.is_superuser
        assert len(password) == 10 and user.check_password(password)
        assert password in response.content.decode()
    event = AuditEvent.objects.get(action="import.create_callers")
    assert all(password not in str(event.details) for _, _, password in created)

    batch.refresh_from_db()
    remaining = {u["name"] for u in services.unmatched_callers(batch)}
    assert remaining == {"Unknown Person"}
    variant = batch.rows.get(row_number=SYNTHETIC_FACTS["caller_spelling_variant"])
    assert variant.values["caller_id"] == User.objects.get(username="synth.caller.one").pk


def test_create_callers_is_all_or_nothing(coordinator, staged_without_callers, make_user):
    make_user("taken.name")
    items = services.unmatched_callers(staged_without_callers)
    entries = [(items[0]["name"], items[0]["username"]), (items[1]["name"], "taken.name")]
    with pytest.raises(services.ImportProblem, match="already taken"):
        services.create_callers_from_workbook(staged_without_callers, entries, coordinator)
    assert not User.objects.filter(username=items[0]["username"]).exists()


def test_create_callers_refuses_names_that_already_match(coordinator, staged_without_callers, make_user):
    make_user("one", ["Caller"], first_name="Synth Caller", last_name="One")
    with pytest.raises(services.ImportProblem):
        services.create_callers_from_workbook(
            staged_without_callers, [("Synth Caller One", "dup.one")], coordinator
        )
    assert not User.objects.filter(username="dup.one").exists()


def test_callers_cannot_create_accounts_from_workbook(client, caller, staged_without_callers):
    client.force_login(caller)
    response = client.post(
        reverse("imports:create_callers", args=[staged_without_callers.pk]),
        {"create": ["0"], "name_0": "Synth Caller One", "username_0": "sneaky"},
    )
    assert response.status_code == 403
    assert not User.objects.filter(username="sneaky").exists()


def test_generated_passwords_are_readable_and_strong():
    from apps.accounts.team import generate_password, suggest_username

    passwords = {generate_password() for _ in range(200)}
    assert len(passwords) == 200
    for p in passwords:
        assert len(p) == 10 and not set(p) & set("0O1lI")
        assert any(c.isdigit() for c in p) and any(c.isupper() for c in p) and any(c.islower() for c in p)
    taken = {"shalini.ji"}
    assert suggest_username("Shalini ji", taken) == "shalini.ji2"
    assert suggest_username("  Vansh  Ji ", taken) == "vansh.ji"


def test_review_summary_counts_problems_and_previews_result(staged):
    review = services.review_summary(staged)
    assert review["total"] == 120
    assert review["needs"] > 0 and review["issues"]["invalid_phones"] == 1
    assert review["issues"]["duplicates"] >= 1 and review["issues"]["missing_fields"] >= 1
    assert review["issues"]["unmatched_callers"] >= 2  # unknown and blank caller rows
    assert review["to_create"] == review["ready"] + review["warning"]
    assert sum(n for _, n in review["per_caller"]) == review["to_allocate"]


# --- Choosing the Saints sheet -----------------------------------------------------


def renamed_workbook(content, settings):
    """The synthetic workbook with a year in the Saints sheet name and extra linked sheets."""
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content))
    book["EMS STALL SEWA"].title = "2026 EMS STALL SEWA"
    summary = book.create_sheet("Area summary", 0)  # first in the file, like a dashboard
    summary.append(["NAME", "AREA", "MOBILE NO"])
    summary.append(["='2026 EMS STALL SEWA'!B2", "Linked", "98765 43210"])
    big = book.create_sheet("Old data")
    for number in range(settings.IMPORT_MAX_ROWS + 50):
        big.append([number])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_finds_saints_sheet_with_year_and_reads_no_other_sheet(workbook, coordinator, settings):
    content = renamed_workbook(workbook[0], settings)
    batch = services.stage_upload(upload(content), coordinator)
    assert batch.contact_sheet == "2026 EMS STALL SEWA"
    assert batch.rows.count() == len(workbook[1])  # nothing from 'Area summary' or 'Old data'
    assert "Area summary" in batch.sheet_names and "Old data" in batch.sheet_names


def test_sheet_name_matching():
    names = ["Area summary", "2026 EMS STALL SEWA", "2026 Calling Sewa status"]
    assert parsing.find_sheet(names, "ems stall sewa") == "2026 EMS STALL SEWA"
    assert parsing.find_sheet(names, "2026 ems  stall sewa", partial=False) == "2026 EMS STALL SEWA"
    assert parsing.find_sheet(names, "stall sewa", partial=False) is None
    assert parsing.find_sheet(["EMS STALL SEWA", "EMS STALL SEWA old"], "ems stall sewa") == "EMS STALL SEWA"
    assert parsing.find_sheet(["2025 EMS STALL SEWA", "2026 EMS STALL SEWA"], "ems stall sewa") is None


def test_typed_sheet_name_must_exist(workbook, coordinator):
    with pytest.raises(services.ImportProblem, match="Sheets in the file"):
        services.stage_upload(upload(workbook[0]), coordinator, contact_sheet="Nope")
    batch = services.stage_upload(upload(workbook[0]), coordinator, contact_sheet="ems stall sewa")
    assert batch.contact_sheet == "EMS STALL SEWA"
