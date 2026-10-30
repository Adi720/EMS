import io
from datetime import date

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.calling.models import CallTask, Outcome
from apps.contacts.models import ContactStatus, StatusReason
from apps.contacts.services import change_status
from apps.reports import metrics


@pytest.fixture
def scenario(make_contact, caller, caller2, allocate, make_session, record, coordinator):
    """caller: 6 contacts, caller2: 3 contacts, with known outcomes (hand-computed below)."""
    first = [make_contact(full_name=f"A{i}", area="North") for i in range(6)]
    second = [make_contact(full_name=f"B{i}", area="South") for i in range(3)]
    for contact in first:
        allocate(contact, caller)
    for contact in second:
        allocate(contact, caller2)
    session = make_session()
    task = {t.contact.full_name: t for t in session.tasks.select_related("contact")}
    record(task["A0"], Outcome.YES)
    record(task["A1"], Outcome.NOT_CONNECTED)
    record(task["A1"], Outcome.WILL_TRY)  # latest: Will Try
    record(task["A2"], Outcome.YES)
    record(task["A2"], Outcome.NOT_CONNECTED)  # latest: Not connected → Unknown
    record(task["A3"], Outcome.NOT_POSSIBLE)
    # A4 not attempted; A5 withdrawn (do not call before calling)
    change_status(
        task["A5"].contact,
        ContactStatus.DO_NOT_CALL,
        coordinator,
        reason_category=StatusReason.REQUESTED_NO_CALLS,
    )
    record(task["B0"], Outcome.WILL_TRY)
    record(task["B1"], Outcome.NOT_CONNECTED)
    return session


EXPECTED_TOTAL = {
    "assigned": 8,
    "completed": 6,
    "not_attempted": 2,
    "connected": 4,
    "yes": 1,
    "not_possible": 1,
    "will_try": 2,
    "unknown": 4,
}
EXPECTED_CALLER = {
    "assigned": 5,
    "completed": 4,
    "not_attempted": 1,
    "connected": 3,
    "yes": 1,
    "not_possible": 1,
    "will_try": 1,
    "unknown": 2,
}
EXPECTED_CALLER2 = {
    "assigned": 3,
    "completed": 2,
    "not_attempted": 1,
    "connected": 1,
    "yes": 0,
    "not_possible": 0,
    "will_try": 1,
    "unknown": 2,
}


def subset(row, keys):
    return {k: row[k] for k in keys}


def test_session_totals_match_hand_computed_values(scenario):
    summary = metrics.summarize(CallTask.objects.filter(session=scenario))
    assert subset(summary, EXPECTED_TOTAL) == EXPECTED_TOTAL
    assert summary["reconciles"]
    assert summary["completion_rate"] == 75  # 6 / 8


def test_per_caller_and_reconciliation(scenario, caller, caller2):
    rows = {r["caller_obj"]: r for r in metrics.by_caller(CallTask.objects.filter(session=scenario))}
    assert subset(rows[caller], EXPECTED_CALLER) == EXPECTED_CALLER
    assert subset(rows[caller2], EXPECTED_CALLER2) == EXPECTED_CALLER2
    for row in rows.values():
        assert row["yes"] + row["not_possible"] + row["will_try"] + row["unknown"] == row["assigned"]
        assert row["completed"] + row["not_attempted"] == row["assigned"]


def test_each_task_falls_in_exactly_one_category(scenario):
    tasks = metrics.assigned_tasks(CallTask.objects.filter(session=scenario))
    categories = [t.report_category for t in tasks]
    summary = metrics.summarize(CallTask.objects.filter(session=scenario))
    for category in ["yes", "not_possible", "will_try", "unknown"]:
        assert categories.count(category) == summary[category]


def test_spec_example_five_attempted_is_five_completed(make_contact, caller, allocate, make_session, record):
    contacts = [make_contact() for _ in range(5)]
    for contact in contacts:
        allocate(contact, caller)
    session = make_session()
    for task, outcome in zip(
        session.tasks.all(), [Outcome.NOT_CONNECTED] * 3 + [Outcome.YES, Outcome.WILL_TRY], strict=True
    ):
        record(task, outcome)
    summary = metrics.summarize(session.tasks.all())
    assert (summary["assigned"], summary["completed"], summary["not_attempted"]) == (5, 5, 0)
    assert summary["yes"] == 1  # Will Try is not confirmed
    assert summary["connected"] == 2


def test_coordinator_dashboard_shows_live_numbers(client, scenario, coordinator, caller):
    client.force_login(coordinator)
    response = client.get(reverse("core:coordinator_dashboard"), {"session": scenario.pk})
    assert subset(response.context["summary"], EXPECTED_TOTAL) == EXPECTED_TOTAL
    response = client.get(
        reverse("core:coordinator_dashboard"), {"session": scenario.pk, "caller": caller.pk}
    )
    assert subset(response.context["summary"], EXPECTED_CALLER) == EXPECTED_CALLER
    assert response.context["history"][0]["session"] == scenario.pk


def test_session_exports(client, scenario, coordinator):
    client.force_login(coordinator)
    url = reverse("reports:session_export", args=[scenario.pk])
    body = client.get(url, {"kind": "summary", "format": "csv"}).content.decode("utf-8-sig")
    lines = body.strip().splitlines()
    assert lines[0].startswith("Caller,Assigned,Completed (attempted)")
    assert lines[-1].startswith("TOTAL,8,6,2,4,1,1,2,4")

    response = client.get(url, {"kind": "tasks", "format": "xlsx"})
    sheet = load_workbook(io.BytesIO(response.content)).active
    names = [row[0].value for row in sheet.iter_rows(min_row=2)]
    assert "A5" not in names  # do-not-call excluded by default
    assert len(names) == 8
    assert "+91" not in str([[c.value for c in row] for row in sheet.iter_rows()])  # no phone numbers
    assert AuditEvent.objects.filter(action__startswith="export.session").count() == 2


def test_monthly_report_sums_sessions(client, scenario, make_session, coordinator, record):
    later = make_session(session_date=date(2026, 11, 8))
    record(later.tasks.first(), Outcome.YES)
    client.force_login(coordinator)
    response = client.get(reverse("reports:monthly"), {"month": "2026-11"})
    summary = response.context["summary"]
    assert summary["assigned"] == 8 + later.tasks.count()
    assert len(response.context["by_session"]) == 2
    assert summary["reconciles"]
    export = client.get(reverse("reports:monthly_export"), {"month": "2026-11", "format": "csv"})
    assert export.status_code == 200


def test_reports_are_coordinator_only(client, scenario, caller):
    client.force_login(caller)
    for url in [
        reverse("core:coordinator_dashboard"),
        reverse("reports:monthly"),
        reverse("reports:monthly_export"),
        reverse("reports:session_export", args=[scenario.pk]) + "?kind=tasks",
        reverse("contacts:export"),
        reverse("audit:event_list"),
    ]:
        assert client.get(url).status_code == 403, url
    assert not AuditEvent.objects.filter(action__startswith="export").exists()
